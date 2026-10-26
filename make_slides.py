"""
產生「解憂雜貨店 (Worry Grocer)」高品質 16:9 PowerPoint 簡報 (.pptx)
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 設定 16:9 寬螢幕比例 (13.333 x 7.5 英吋)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 配色定義
    COLOR_BG = RGBColor(250, 247, 242)       # #FAF7F2 暖奶油底色
    COLOR_INK = RGBColor(44, 38, 35)          # #2C2623 沉穩黑墨
    COLOR_ACCENT = RGBColor(200, 90, 50)      # #C85A32 磚紅焦點色
    COLOR_MUTED = RGBColor(124, 110, 101)     # #7C6E65 柔和灰褐
    COLOR_WHITE = RGBColor(255, 255, 255)     # #FFFFFF 卡片白
    COLOR_BORDER = RGBColor(232, 223, 213)    # #E8DFD5 邊框灰

    def add_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category="期中專題 · 11/03 成果分享"):
        # 分類標籤
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(10), Inches(0.4))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = f"🏮 解憂雜貨店 (Worry Grocer)  |  {category}"
        p_c.font.size = Pt(12)
        p_c.font.bold = True
        p_c.font.color.rgb = COLOR_ACCENT

        # 主標題
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(28)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_INK

    blank_layout = prs.slide_layouts[6]

    # -------------------------------------------------------------
    # Slide 1: 封面
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    add_slide_background(s1)

    # 裝飾卡片
    card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.0), Inches(10.333), Inches(5.5))
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_WHITE
    card.line.color.rgb = COLOR_BORDER
    card.line.width = Pt(1.5)

    tb = s1.shapes.add_textbox(Inches(2.0), Inches(1.4), Inches(9.333), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.alignment = PP_ALIGN.CENTER
    p0.text = "✨ 用 AI 做出一件能分享給別人的點子"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ACCENT
    p0.space_after = Pt(20)

    p1 = tf.add_paragraph()
    p1.alignment = PP_ALIGN.CENTER
    p1.text = "🏮 解憂雜貨店 (Worry Grocer)"
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_INK
    p1.space_after = Pt(12)

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "你的深夜心事樹洞與心靈處方箋生成器"
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_ACCENT
    p2.space_after = Pt(24)

    p3 = tf.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    p3.text = "「把你的煩惱投進信箱，明天太陽依舊升起。」\n結合多重人格陪伴、非暴力溝通與拍立得微行動小卡"
    p3.font.size = Pt(14)
    p3.font.color.rgb = COLOR_MUTED
    p3.space_after = Pt(30)

    p4 = tf.add_paragraph()
    p4.alignment = PP_ALIGN.CENTER
    p4.text = "期中專題成果發表 · 11/03 課堂分享"
    p4.font.size = Pt(13)
    p4.font.color.rgb = COLOR_MUTED

    # -------------------------------------------------------------
    # Slide 2: 痛點與為什麼做這個 (Problem & Insight)
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_background(s2)
    add_header(s2, "我們都有想傾訴的心事，卻不知向誰開口")

    cols = [
        ("🤐 痛點一：想傾訴，卻怕給朋友添麻煩", "深夜焦慮或課業挫折，總覺得發在社群太負面、跟朋友抱怨又怕打擾別人生活，只能獨自悶在心裡無效內耗。"),
        ("🤖 痛點二：傳統 AI 對話過於冰冷說教", "常見的聊天機器人只會條列式給大道理或 SOP 建議，缺乏手寫信的溫度、儀式感與真實的情感接納與共鳴。"),
        ("📤 痛點三：缺乏可隨手分享的形式", "聊天截圖隱私度高且不美觀。我們需要一種去識別化、精緻如拍立得的美感產出，把溫暖與力量自然傳遞出去。")
    ]

    for i, (title, desc) in enumerate(cols):
        left = Inches(0.8 + i * 4.0)
        c = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.9), Inches(3.7), Inches(4.8))
        c.fill.solid()
        c.fill.fore_color.rgb = COLOR_WHITE
        c.line.color.rgb = COLOR_BORDER

        tb = s2.shapes.add_textbox(left + Inches(0.25), Inches(2.2), Inches(3.2), Inches(4.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_INK
        p_t.space_after = Pt(16)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = COLOR_MUTED

    # -------------------------------------------------------------
    # Slide 3: 核心體驗與亮點 (Core Experience)
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_background(s3)
    add_header(s3, "從寫信到開立處方箋的儀式感")

    features = [
        ("🌿 三大筆友人格自由挑選", "• 溫柔木木：治癒陪伴、深層共情、無條件包容脆弱\n• 毒舌阿柴：幽默吐槽、一針見血戳破內耗、拉回現實\n• 哲學爺爺：宏觀視角、以大自然與歲月為隱喻開解"),
        ("📮 沉浸式書信儀式動效", "• 日系橫線信紙親手書寫，自訂匿名稱呼\n• 投遞時蓋上專屬日期郵戳，郵差送信小動效\n• 拆信展開時伴隨慶祝紙屑，重溫期待回信的浪漫"),
        ("📸 今日心靈處方箋拍立得", "• 直擊心靈的治癒金句 + 今天就能做的一件微行動任務\n• 瀏覽器直接渲染生成高解析度 PNG 圖片\n• 一鍵下載，輕鬆分享到 IG 限動或收藏")
    ]

    for i, (f_title, f_desc) in enumerate(features):
        top = Inches(1.8 + i * 1.7)
        c = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(11.7), Inches(1.5))
        c.fill.solid()
        c.fill.fore_color.rgb = COLOR_WHITE
        c.line.color.rgb = COLOR_BORDER

        tb = s3.shapes.add_textbox(Inches(1.1), top + Inches(0.15), Inches(11.1), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = f_title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_ACCENT
        p_t.space_after = Pt(6)

        p_d = tf.add_paragraph()
        p_d.text = f_desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_MUTED

    # -------------------------------------------------------------
    # Slide 4: 怎麼跟 AI 一起做 (AI Co-Creation)
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_background(s4)
    add_header(s4, "人機協作：如何調教出有溫度的 AI？")

    ai_steps = [
        ("1. Prompt Engineering 導入非暴力溝通", "嚴格設計 System Instruction，結合心理學 NVC 框架：觀察 ➔ 確認感受 ➔ 共情 ➔ 給予微行動。嚴禁冷冰冰說教，賦予三大店長鮮活生命力。"),
        ("2. 結構化輸出 (JSON Mode)", "讓 Gemini API 輸出包含 letter (長信)、quote (金句)、action (微行動) 的 JSON 格式，方便前端元件精準排版至信紙與拍立得。"),
        ("3. 客戶端渲染攻堅 (html2canvas)", "與 AI 協同除錯，在不依賴後端伺服器的情況下調校 Retina 2.5x 解析度與 CSS 盒模型，純瀏覽器端直接匯出無損拍立得圖片。"),
        ("4. 雙模設計 (高可用離線備援庫)", "設計了免 API Key 的智慧離線備援庫，根據關鍵字智慧匹配情境回信，確保課堂 Demo 與公開分享時 100% 順暢不卡關。")
    ]

    for i, (st, sd) in enumerate(ai_steps):
        row = i // 2
        col = i % 2
        left = Inches(0.8 + col * 6.0)
        top = Inches(1.9 + row * 2.5)

        c = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.7), Inches(2.2))
        c.fill.solid()
        c.fill.fore_color.rgb = COLOR_WHITE
        c.line.color.rgb = COLOR_BORDER

        tb = s4.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), Inches(5.3), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = st
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_INK
        p_t.space_after = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = sd
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_MUTED

    # -------------------------------------------------------------
    # Slide 5: 未來展望與擴充 (Roadmap & Impact)
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_background(s5)
    add_header(s5, "下一步展望：一個真正能分享給身邊人的點子")

    roadmaps = [
        ("01", "🌌 漂流星空共鳴池", "經寄信者同意後匿名投遞到公開星空，訪客可以在夜空中點擊星點閱讀回信，並給予無聲的「溫暖擁抱」，形成互助社群。"),
        ("02", "🎙️ 多模態語音讀信 (TTS)", "結合語音合成技術，讓溫柔木木或哲學爺爺用溫暖且具呼吸感的聲線親口為你朗讀信件，在深夜戴上耳機感受聲音的陪伴。"),
        ("03", "🛡️ 心理危機防護與安心轉介", "導入關鍵字敏感度篩檢，若偵測到重度心理創傷或自我傷害傾向，回信中自動溫柔引導校園心輔中心與衛福部 1925 安心專線。")
    ]

    for i, (num, rt, rd) in enumerate(roadmaps):
        left = Inches(0.8 + i * 4.0)
        c = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.9), Inches(3.7), Inches(4.8))
        c.fill.solid()
        c.fill.fore_color.rgb = COLOR_WHITE
        c.line.color.rgb = COLOR_BORDER

        tb = s5.shapes.add_textbox(left + Inches(0.25), Inches(2.2), Inches(3.2), Inches(4.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p_n = tf.paragraphs[0]
        p_n.text = num
        p_n.font.size = Pt(28)
        p_n.font.bold = True
        p_n.font.color.rgb = COLOR_ACCENT
        p_n.space_after = Pt(8)

        p_t = tf.add_paragraph()
        p_t.text = rt
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_INK
        p_t.space_after = Pt(14)

        p_d = tf.add_paragraph()
        p_d.text = rd
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_MUTED

    # -------------------------------------------------------------
    # Slide 6: 結語與現場 Demo 邀請
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_background(s6)

    c6 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.0), Inches(10.333), Inches(5.5))
    c6.fill.solid()
    c6.fill.fore_color.rgb = COLOR_WHITE
    c6.line.color.rgb = COLOR_BORDER

    tb6 = s6.shapes.add_textbox(Inches(2.0), Inches(1.4), Inches(9.333), Inches(4.7))
    tf6 = tb6.text_frame
    tf6.word_wrap = True

    p6_0 = tf6.paragraphs[0]
    p6_0.alignment = PP_ALIGN.CENTER
    p6_0.text = "🏮 歡迎現場體驗"
    p6_0.font.size = Pt(16)
    p6_0.font.bold = True
    p6_0.font.color.rgb = COLOR_ACCENT
    p6_0.space_after = Pt(10)

    p6_1 = tf6.add_paragraph()
    p6_1.alignment = PP_ALIGN.CENTER
    p6_1.text = "把心事交給雜貨店，明天又是新的一天"
    p6_1.font.size = Pt(32)
    p6_1.font.bold = True
    p6_1.font.color.rgb = COLOR_INK
    p6_1.space_after = Pt(20)

    p6_2 = tf6.add_paragraph()
    p6_2.alignment = PP_ALIGN.CENTER
    p6_2.text = "線上體驗網址（GitHub Pages）：\nhttps://411411324.github.io/worry-grocer/\n\n專案開源 Repository：\nhttps://github.com/411411324/worry-grocer"
    p6_2.font.size = Pt(14)
    p6_2.font.bold = True
    p6_2.font.color.rgb = COLOR_ACCENT
    p6_2.space_after = Pt(24)

    p6_3 = tf6.add_paragraph()
    p6_3.alignment = PP_ALIGN.CENTER
    p6_3.text = "謝謝老師與全班同學！歡迎大家提問與交流 ❤️"
    p6_3.font.size = Pt(14)
    p6_3.font.color.rgb = COLOR_MUTED

    # 輸出儲存
    output_path = r"C:\Users\sfyeh\.gemini\antigravity\scratch\worry-grocer\解憂雜貨店_期中簡報.pptx"
    prs.save(output_path)
    print(f"Presentation saved to: {output_path}")

if __name__ == "__main__":
    create_presentation()
